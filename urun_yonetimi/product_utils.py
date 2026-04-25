import customtkinter as ctk
from tkinter import ttk, messagebox, filedialog, simpledialog
import openpyxl
from core.api_client import ApiClientError, delete_products
from core.database import veritabani_baglanti
from core.session import get_app_token
from urun_detay.product_detail import urun_detay_karti
from urun_yonetimi.product_tree import urun_agaci_ekrani
from maliyet.cost_calculator import maliyet_hesapla
import threading


KOLONLAR = [
    ("urun_kodu", "Ürün Kodu", "entry"), ("urun_adi", "Ürün Adı", "entry"),
    ("aciklama", "Açıklama", "entry"), ("urun_kategorisi", "Kategori", "combobox"),
    ("urun_tipi", "Ürün Tipi", "combobox"), ("urun_modeli", "Model", "entry"),
    ("maliyet", "Genel Toplam Maliyet", "nosrch"), ("filtre_medyasi", "Filtre Medyası", "combobox"),
    ("filtre_medyasi_kodu", "Filtre Medyası Kodu", "entry"), ("patlac_kumanda_tipi", "Patlaç Kontrol", "combobox"),
    ("toplam_filtre_alani", "Toplam Filtre Alanı", "entry"), ("debi", "Debi", "entry"),
    ("fan_basinc", "Basınç", "entry"), ("fan_basinc_birimi", "Basınç Birimi", "combobox"),
    ("motor", "Motor", "entry"), ("fan_kumanda_tipi", "Fan Pano Tipi", "combobox"),
    ("patlama_kapagi", "Patlama Kapağı", "entry"), ("filtre_elemani_sayisi", "Filtre Sayısı", "entry"),
]

def filtre_seceneklerini_yukle(pencere, filtre_opsiyonlari, filtre_widgets):
    """Filtre seçeneklerini asenkron olarak yükler"""
    db = None
    try:
        db = veritabani_baglanti()
        cursor = db.cursor(buffered=True)
        
        # Her kolon için ayrı ayrı filtre seçeneklerini çek
        for kolon, _, _ in KOLONLAR:
            if kolon != "maliyet":  # Maliyet kolonu için filtre yok
                try:
                    cursor.execute(f"SELECT DISTINCT {kolon} FROM urunler WHERE {kolon} IS NOT NULL AND {kolon} != '' ORDER BY {kolon}")
                    filtre_opsiyonlari[kolon] = [row[0] for row in cursor.fetchall()]
                except Exception as e:
                    print(f"Kolon {kolon} için filtre seçenekleri yüklenirken hata: {e}")
                    filtre_opsiyonlari[kolon] = []
        
        # UI'ı güncelle
        pencere.after(0, lambda: filtre_ui_guncelle(filtre_widgets, filtre_opsiyonlari))
        
    except Exception as e:
        print(f"Filtre seçenekleri yükleme hatası: {e}")
    finally:
        if db and db.is_connected():
            db.close()

def filtre_ui_guncelle(filtre_widgets, filtre_opsiyonlari):
    """Filtre UI'ını günceller"""
    for kolon, widget in filtre_widgets.items():
        if kolon in filtre_opsiyonlari:
            widget.configure(values=filtre_opsiyonlari[kolon])

def veri_filtrele(veriler, arama_var, aktif_filtreler):
    """Verileri filtreler"""
    sonuc = []
    arama_terimi = arama_var.get().lower()
    
    # Sütun indekslerini tanımla (veritabanındaki sırayla)
    kolon_indeksleri = {
        "urun_kodu": 1,
        "urun_adi": 2,
        "aciklama": 3,
        "urun_kategorisi": 4,
        "urun_tipi": 5,
        "urun_modeli": 6,
        "maliyet": 7,
        "filtre_medyasi": 8,
        "filtre_medyasi_kodu": 9,
        "patlac_kumanda_tipi": 10,
        "toplam_filtre_alani": 11,
        "debi": 12,
        "fan_basinc": 13,
        "fan_basinc_birimi": 14,
        "motor": 15,
        "fan_kumanda_tipi": 16,
        "patlama_kapagi": 17,
        "filtre_elemani_sayisi": 18
    }
    
    for row in veriler:
        uygun = True
        
        # Genel arama kontrolü
        if arama_terimi:
            arama_bulundu = False
            for kolon, _, _ in KOLONLAR:
                if kolon == "id":
                    continue
                kolon_index = kolon_indeksleri.get(kolon)
                if kolon_index is not None and kolon_index < len(row):
                    deger = str(row[kolon_index]) if row[kolon_index] is not None else ""
                    if arama_terimi in deger.lower():
                        arama_bulundu = True
                        break
            if not arama_bulundu:
                uygun = False
        
        # Aktif filtreler kontrolü
        for kolon, filtre_deger in aktif_filtreler.items():
            kolon_index = kolon_indeksleri.get(kolon)
            if kolon_index is not None and kolon_index < len(row):
                deger = str(row[kolon_index]) if row[kolon_index] is not None else ""
                if filtre_deger.lower() not in deger.lower():
                    uygun = False
                    break
            else:
                uygun = False
                break
        
        if uygun:
            sonuc.append(row)
    
    return sonuc

def veri_yukle_async(pencere, progress_bar, progress_label, tablo_ui_guncelle_callback):
    """Verileri asenkron olarak yükler - Genel kullanım için tüm ürünleri çeker"""
    db = None
    try:
        progress_bar.set(0.3)
        progress_label.configure(text="Veritabanından veriler çekiliyor...")
        
        db = veritabani_baglanti()
        cursor = db.cursor(buffered=True)
        
        # Tüm ürünleri çek (genel kullanım için)
        cursor.execute(
            "SELECT id, urun_kodu, urun_adi, aciklama, urun_kategorisi, urun_tipi, urun_modeli, maliyet, "
            "filtre_medyasi, filtre_medyasi_kodu, patlac_kumanda_tipi, toplam_filtre_alani, debi, fan_basinc, "
            "fan_basinc_birimi, motor, fan_kumanda_tipi, patlama_kapagi, filtre_elemani_sayisi, "
            "kanal_capi, kanal_boyu, kanal_et_kalinlik, flans_capi, flans_kalinlik, "
            "malzeme_maliyeti, iscilik_maliyeti, uretim_gideri, yonetim_gideri, alt_urun_maliyeti, maliyet_hesaplama_tarihi "
            "FROM urunler"
        )
        
        progress_bar.set(0.7)
        progress_label.configure(text="Veriler işleniyor...")
        
        veriler = cursor.fetchall()
        
        # Global değişken olarak kaydet
        import sys
        sys.urunler_veritabani = veriler
        
        progress_bar.set(1.0)
        progress_label.configure(text="Veriler yüklendi!")
        
        # UI güncellemesini ana thread'de yap
        pencere.after(0, lambda: tablo_ui_guncelle_callback(veriler))
        
    except Exception as e:
        print(f"Veri yükleme hatası: {e}")
        import traceback
        traceback.print_exc()
        pencere.after(0, lambda: tablo_ui_guncelle_callback([]))
    finally:
        if db and db.is_connected():
            db.close()

def disari_aktar(tree):
    """Dışa aktarma fonksiyonu - Treeview'dan doğrudan veri alır"""
    try:
        # Treeview'dan tüm verileri al
        tree_items = tree.get_children()
        
        if not tree_items:
            messagebox.showwarning("Uyarı", "Dışa aktarılacak veri bulunamadı.")
            return
        
        dosya_yolu = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel dosyası", "*.xlsx")],
            title="Dosyayı Kaydet"
        )
        if not dosya_yolu:
            return
        
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Ürünler"
        
        # Başlıkları ekle
        basliklar = [k[1] for k in KOLONLAR]
        sheet.append(basliklar)
        
        # Treeview'dan verileri al ve Excel'e ekle
        veri_sayisi = 0
        for item in tree_items:
            values = tree.item(item)["values"]
            sheet.append(values)
            veri_sayisi += 1
        
        workbook.save(dosya_yolu)
        messagebox.showinfo("Başarılı", f"Veriler başarıyla dışa aktarıldı:\n{dosya_yolu}\n\nToplam {veri_sayisi} ürün aktarıldı.")
    except Exception as e:
        print(f"Dışa aktarma hatası: {e}")
        messagebox.showerror("Hata", f"Dışa aktarma sırasında bir hata oluştu:\n{e}")

def urun_sil(tree, yenile, pencere):
    """Tek ürün silme fonksiyonu"""
    return akilli_urun_sil(tree, yenile, pencere)


def toplu_urun_sil(tree, yenile, pencere):
    """Toplu ürün silme fonksiyonu"""
    return akilli_urun_sil(tree, yenile, pencere)


def akilli_urun_sil(tree, yenile, pencere):
    """Akıllı ürün silme fonksiyonu - tek veya toplu silme (UI donmasın diye arkaplanda)."""
    selected_items = tree.selection()
    if not selected_items:
        messagebox.showwarning("Uyarı", "Lütfen silinecek ürünü/ürünleri seçin.", parent=pencere)
        return

    # Seçili ürün kodlarını al
    urun_kodlari = []
    for item in selected_items:
        urun_kodu = str(tree.item(item)["values"][0]).strip()
        urun_kodlari.append(urun_kodu)

    # Onay pencereleri ana thread'de
    if len(urun_kodlari) == 1:
        urun_kodu = urun_kodlari[0]
        onay = messagebox.askyesno(
            "Onay",
            f"'{urun_kodu}' kodlu ürünü ve bağlı tüm verilerini (ürün ağacı, işçilik) kalıcı olarak silmek istediğinize emin misiniz?",
            icon='warning',
            parent=pencere,
        )
        if not onay:
            return
    else:
        onay = messagebox.askyesno(
            "Toplu Silme Onayı",
            f"{len(urun_kodlari)} ürünü kalıcı olarak silmek istediğinize emin misiniz?\n\n"
            f"Silinecek ürünler: {', '.join(map(str, urun_kodlari[:5]))}{'...' if len(urun_kodlari) > 5 else ''}",
            icon='warning',
            parent=pencere,
        )
        if not onay:
            return

    # Basit bir ilerleme penceresi oluştur (modal)
    progress_win = ctk.CTkToplevel(pencere)
    progress_win.title("Siliniyor...")
    try:
        progress_win.transient(pencere)
        progress_win.grab_set()
    except Exception:
        pass
    ctk.CTkLabel(progress_win, text="Silme işlemi devam ediyor, lütfen bekleyin...", font=ctk.CTkFont(size=14)).pack(padx=20, pady=(20, 10))
    pb = ctk.CTkProgressBar(progress_win, mode="indeterminate", width=260)
    pb.pack(padx=20, pady=(0, 20))
    try:
        pb.start()
    except Exception:
        pass

    def arkaplanda_sil():
        db = None
        sonuc_mesaji = ""
        hata_mesaji = None
        try:
            app_token = get_app_token()
            if app_token:
                try:
                    response = delete_products(app_token, urun_kodlari) or {}
                    basarili_silinen = int((response or {}).get("deleted_count") or 0)
                    hatali_silinen = int((response or {}).get("blocked_count") or 0)
                    if len(urun_kodlari) == 1:
                        if basarili_silinen == 1 and hatali_silinen == 0:
                            sonuc_mesaji = "Ürün ve bağlı tüm verileri başarıyla silindi."
                        elif hatali_silinen >= 1:
                            sonuc_mesaji = "Ürün silinemedi (proje listesinde kullanılıyor veya hata oluştu)."
                    else:
                        sonuc_mesaji = "Toplu silme tamamlandı:\n"
                        if basarili_silinen > 0:
                            sonuc_mesaji += f"✅ {basarili_silinen} ürün başarıyla silindi\n"
                        if hatali_silinen > 0:
                            sonuc_mesaji += f"âŒ {hatali_silinen} Ã¼rÃ¼n silinemedi (proje listesinde kullanÄ±lÄ±yor veya hata oluÅŸtu)"
                    db = None
                    raise StopIteration
                except ApiClientError as e:
                    hata_mesaji = str(e)
                    raise StopIteration

            db = veritabani_baglanti()
            cursor = db.cursor(buffered=True)

            basarili_silinen = 0
            hatali_silinen = 0

            for urun_kodu in urun_kodlari:
                try:
                    # Ürün ID'sini al
                    cursor.execute("SELECT id FROM urunler WHERE urun_kodu = %s", (urun_kodu,))
                    urun = cursor.fetchone()
                    if not urun:
                        hatali_silinen += 1
                        continue

                    urun_id = urun[0]

                    # Güvenlik kontrolü
                    cursor.execute("SELECT COUNT(*) FROM proje_listesi_icerigi WHERE urun_id = %s", (urun_id,))
                    kullanim_sayisi = cursor.fetchone()[0]
                    if kullanim_sayisi > 0:
                        hatali_silinen += 1
                        continue

                    # Transaction ile güvenli silme (tek tek)
                    db.autocommit = False
                    cursor.execute("DELETE FROM urun_agaci WHERE urun_id = %s", (urun_id,))
                    cursor.execute("DELETE FROM urun_iscilik WHERE urun_id = %s", (urun_id,))
                    cursor.execute("DELETE FROM urunler WHERE id = %s", (urun_id,))
                    db.commit()
                    db.autocommit = True
                    basarili_silinen += 1
                except Exception as e:
                    try:
                        if db:
                            db.rollback()
                            db.autocommit = True
                    except Exception:
                        pass
                    hatali_silinen += 1
                    print(f"Ürün {urun_kodu} silinirken hata: {e}")

            if len(urun_kodlari) == 1:
                if basarili_silinen == 1 and hatali_silinen == 0:
                    sonuc_mesaji = "Ürün ve bağlı tüm verileri başarıyla silindi."
                elif hatali_silinen == 1:
                    sonuc_mesaji = "Ürün silinemedi (proje listesinde kullanılıyor veya hata oluştu)."
            else:
                sonuc_mesaji = "Toplu silme tamamlandı:\n"
                if basarili_silinen > 0:
                    sonuc_mesaji += f"✅ {basarili_silinen} ürün başarıyla silindi\n"
                if hatali_silinen > 0:
                    sonuc_mesaji += f"âŒ {hatali_silinen} Ã¼rÃ¼n silinemedi (proje listesinde kullanÄ±lÄ±yor veya hata oluÅŸtu)"

        except StopIteration:
            pass
        except Exception as e:
            hata_mesaji = str(e)
        finally:
            if db and db.is_connected():
                try:
                    db.autocommit = True
                except Exception:
                    pass
                db.close()

        # UI güncellemeleri ana thread'de
        def tamamla_ui():
            try:
                if pb:
                    pb.stop()
            except Exception:
                pass
            try:
                progress_win.destroy()
            except Exception:
                pass
            try:
                yenile()
            except Exception:
                pass
            if hata_mesaji:
                messagebox.showerror("Hata", f"Silme sırasında bir hata oluştu: {hata_mesaji}", parent=pencere)
            else:
                baslik = "Toplu Silme Sonucu" if len(urun_kodlari) > 1 else "Başarılı"
                messagebox.showinfo(baslik, sonuc_mesaji, parent=pencere)

        try:
            pencere.after(0, tamamla_ui)
        except Exception:
            pass

    threading.Thread(target=arkaplanda_sil, daemon=True).start()

def urun_duzenle(tree, yenile, kullanici_rolu, pencere):
    """Ürün düzenleme fonksiyonu"""
    selected = tree.selection()
    if not selected:
        messagebox.showwarning("Uyarı", "Lütfen düzenlenecek ürünü seçin.", parent=pencere)
        return
    
    # Treeview'da ID yok, ilk kolon ürün kodu
    urun_kodu = str(tree.item(selected[0])["values"][0]).strip()
    db = None
    try:
        db = veritabani_baglanti()
        c = db.cursor()
        # Ürün detay kartı için tüm sütunları çek
        c.execute("SELECT * FROM urunler WHERE urun_kodu = %s", (urun_kodu,))
        
        veri = c.fetchone()
        if veri:
            urun_detay_karti(pencere, veri, duzenleme=True, yenile_fonksiyonu=yenile, kullanici_rolu=kullanici_rolu)
        else:
            messagebox.showerror("Hata", "Ürün bulunamadı!", parent=pencere)
    except Exception as e:
        messagebox.showerror("Hata", f"Ürün düzenleme hatası: {e}", parent=pencere)
    finally:
        if db and db.is_connected():
            db.close()

def urun_agaci(tree, kullanici_rolu, tablo_yenile):
    """Ürün ağacı açma fonksiyonu"""
    selected = tree.selection()
    if not selected:
        messagebox.showwarning("Uyarı", "Lütfen bir ürün seçin.")
        return
    
    # Treeview'da ID yok, ilk kolon ürün kodu
    urun_kodu = str(tree.item(selected[0])["values"][0]).strip()
    db = None
    try:
        db = veritabani_baglanti()
        c = db.cursor()
        c.execute("SELECT id FROM urunler WHERE urun_kodu = %s", (urun_kodu,))
        urun = c.fetchone()
        if urun:
            # Global yenileme fonksiyonunu kullan
            def yenileme_wrapper():
                try:
                    import sys
                    if hasattr(sys, 'urunler_tablo_yenile'):
                        sys.urunler_tablo_yenile()
                    else:
                        tablo_yenile()
                except Exception as e:
                    print(f"Yenileme fonksiyonu hatası: {e}")
            
            urun_agaci_ekrani(urun[0], kullanici_rolu, yenileme_fonksiyonu=yenileme_wrapper)
    except Exception as e:
        messagebox.showerror("Hata", f"Ürün ağacı açılırken bir hata oluştu: {e}")
    finally:
        if db and db.is_connected():
            db.close()

# === YENİ: ÜRÜN KOPYALAMA ===
def urun_kopyala(tree, yenile, pencere):
    """Seçili ürün/ürünleri, ürün ağacı ve işçilikleriyle birlikte kopyalar (UI donmasın diye arkaplanda)."""
    selected = tree.selection()
    if not selected:
        messagebox.showwarning("Uyarı", "Lütfen kopyalanacak ürünü seçin.", parent=pencere)
        return

    # 1) Gerekli tüm kullanıcı girdilerini ÖNCE al (UI thread)
    kopya_listesi = []  # (eski_kod, yeni_kod)
    for item in selected:
        eski_kod = str(tree.item(item)["values"][0]).strip()
        yeni_kod = simpledialog.askstring(
            "Yeni Ürün Kodu",
            f"{eski_kod} için yeni ürün kodunu girin:",
            parent=pencere,
        )
        if yeni_kod is None:
            continue
        yeni_kod = yeni_kod.strip()
        if not yeni_kod:
            messagebox.showwarning("Uyarı", "Ürün kodu boş olamaz.", parent=pencere)
            continue
        kopya_listesi.append((eski_kod, yeni_kod))

    if not kopya_listesi:
        return

    # 2) Progress penceresi
    progress_win = ctk.CTkToplevel(pencere)
    progress_win.title("Kopyalanıyor...")
    try:
        progress_win.transient(pencere)
        progress_win.grab_set()
    except Exception:
        pass
    ctk.CTkLabel(progress_win, text="Ürün(ler) kopyalanıyor, lütfen bekleyin...", font=ctk.CTkFont(size=14)).pack(padx=20, pady=(20, 10))
    pb = ctk.CTkProgressBar(progress_win, mode="indeterminate", width=260)
    pb.pack(padx=20, pady=(0, 20))
    try:
        pb.start()
    except Exception:
        pass

    def arkaplanda_kopyala():
        db = None
        hata_mesaji = None
        sonuc_bilgi = []
        try:
            db = veritabani_baglanti()
            cursor = db.cursor(buffered=True)

            for eski_kod, yeni_kod in kopya_listesi:
                try:
                    # Eski ürün kaydını al ve kolon listesini hazırla
                    cursor.execute("SELECT * FROM urunler WHERE urun_kodu = %s", (eski_kod,))
                    eski_veri = cursor.fetchone()
                    if not eski_veri:
                        sonuc_bilgi.append(f"âŒ {eski_kod}: Ã¼rÃ¼n bulunamadÄ±")
                        continue
                    kolonlar = [desc[0] for desc in cursor.description]

                    # Aynı kod var mı kontrol
                    cursor.execute("SELECT COUNT(*) FROM urunler WHERE urun_kodu = %s", (yeni_kod,))
                    if cursor.fetchone()[0] > 0:
                        sonuc_bilgi.append(f"âŒ {yeni_kod}: kod zaten mevcut")
                        continue

                    # Transaction - her ürün için ayrı
                    try:
                        db.autocommit = False
                    except Exception:
                        pass

                    # ÜRÜNLER'e ekle
                    insert_cols = [k for k in kolonlar if k not in ("id", "urun_kodu")]
                    insert_vals = [eski_veri[kolonlar.index(col)] for col in insert_cols]
                    placeholders = ", ".join(["%s"] * len(insert_cols))
                    cursor.execute(
                        f"INSERT INTO urunler (urun_kodu, {', '.join(insert_cols)}) VALUES (%s, {placeholders})",
                        (yeni_kod, *insert_vals),
                    )
                    yeni_id = cursor.lastrowid

                    # ÃœRÃœN AÄACI
                    cursor.execute(
                        "SELECT malzeme_kodu, malzeme_adi, miktar, malzeme_tipi, alt_urun_id FROM urun_agaci WHERE urun_id = %s",
                        (eski_veri[0],),
                    )
                    urun_agaci_kayitlar = cursor.fetchall()
                    if urun_agaci_kayitlar:
                        cursor.executemany(
                            "INSERT INTO urun_agaci (urun_id, malzeme_kodu, malzeme_adi, miktar, malzeme_tipi, alt_urun_id) VALUES (%s,%s,%s,%s,%s,%s)",
                            [(yeni_id, *row) for row in urun_agaci_kayitlar],
                        )

                    # Ä°ÅÃ‡Ä°LÄ°K
                    cursor.execute("SELECT iscilik_tipi, usta_saat, yardimci_saat FROM urun_iscilik WHERE urun_id = %s", (eski_veri[0],))
                    iscilik_kayitlar = cursor.fetchall()
                    if iscilik_kayitlar:
                        cursor.executemany(
                            "INSERT INTO urun_iscilik (urun_id, iscilik_tipi, usta_saat, yardimci_saat) VALUES (%s,%s,%s,%s)",
                            [(yeni_id, *row) for row in iscilik_kayitlar],
                        )

                    # Maliyet hesapla
                    dict_cursor = db.cursor(dictionary=True, buffered=True)
                    maliyet_hesapla(yeni_id, dict_cursor)

                    db.commit()
                    try:
                        db.autocommit = True
                    except Exception:
                        pass
                    sonuc_bilgi.append(f"✅ {eski_kod} -> {yeni_kod} kopyalandı")
                except Exception as e:
                    try:
                        if db:
                            db.rollback()
                            db.autocommit = True
                    except Exception:
                        pass
                    sonuc_bilgi.append(f"âŒ {eski_kod} -> {yeni_kod}: {e}")

        except Exception as e:
            hata_mesaji = str(e)
        finally:
            if db and db.is_connected():
                try:
                    db.autocommit = True
                except Exception:
                    pass
                db.close()

        def tamamla_ui():
            try:
                if pb:
                    pb.stop()
            except Exception:
                pass
            try:
                progress_win.destroy()
            except Exception:
                pass
            try:
                yenile()
            except Exception:
                pass
            if hata_mesaji:
                messagebox.showerror("Hata", f"Kopyalama sırasında hata oluştu: {hata_mesaji}", parent=pencere)
            else:
                ozet = "\n".join(sonuc_bilgi) if sonuc_bilgi else "İşlem tamamlandı."
                messagebox.showinfo("Kopyalama Sonucu", ozet, parent=pencere)

        try:
            pencere.after(0, tamamla_ui)
        except Exception:
            pass

    threading.Thread(target=arkaplanda_kopyala, daemon=True).start()

def urun_detayini_getir(event, kullanici_rolu, pencere):
    """Ürün detayını getirme fonksiyonu"""
    tree_widget = event.widget
    selected = tree_widget.selection()
    if selected:
        # Treeview'da ID yok, ilk kolon ürün kodu
        urun_kodu = str(tree_widget.item(selected[0])["values"][0]).strip()
        db = None
        try:
            db = veritabani_baglanti()
            c = db.cursor()
            # Ürün detay kartı için tüm sütunları çek
            c.execute("SELECT * FROM urunler WHERE urun_kodu = %s", (urun_kodu,))
            
            veri = c.fetchone()
            if veri:
                urun_detay_karti(pencere, veri, duzenleme=False, kullanici_rolu=kullanici_rolu)
            else:
                messagebox.showerror("Hata", "Ürün bulunamadı!", parent=pencere)
        except Exception as e:
            messagebox.showerror("Hata", f"Ürün detayı getirilirken hata oluştu: {e}", parent=pencere)
        finally:
            if db and db.is_connected():
                db.close() 
